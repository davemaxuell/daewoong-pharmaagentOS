from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any
from urllib.parse import urlencode, urljoin

from bs4 import BeautifulSoup
from openpyxl import load_workbook

from app.ingestion import CandidateMetadata
from app.parsing import normalize_whitespace
from app.security.urls import UnsafeUrlError, canonicalize_fda_url


class ListingDiscoveryError(ValueError):
    """Raised when an FDA listing cannot be represented without guessing."""


DATE_FORMATS = ("%m/%d/%Y", "%Y-%m-%d", "%B %d, %Y", "%b %d, %Y")


def _column_key(value: object) -> str:
    text = normalize_whitespace(str(value or "")).casefold()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


COLUMN_ALIASES = {
    "posted_date": {"posted date", "date posted", "posted"},
    "issue_date": {"letter issue date", "issue date", "date issued"},
    "company_name": {"company name", "company", "firm name", "firm"},
    "country": {"country", "country name"},
    "issuing_office": {"issuing office", "office", "center"},
    "subject": {"subject", "letter subject"},
    "canonical_url": {
        "letter url",
        "warning letter url",
        "url",
        "letter link",
        "company url",
    },
    "marcs_cms_number": {"marcs cms number", "marcs cms", "cms number"},
    "fda_reference_number": {"reference number", "fda reference number"},
    "response_url": {"response letter", "response url", "response letter url"},
    "closeout_url": {
        "close out letter",
        "closeout letter",
        "close out url",
        "closeout url",
    },
}


def _canonical_column(value: object) -> str:
    key = _column_key(value)
    for canonical, aliases in COLUMN_ALIASES.items():
        if key in aliases:
            return canonical
    return key.replace(" ", "_")


def _parse_date(value: object) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = normalize_whitespace(str(value))
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


@dataclass(frozen=True)
class ListingCandidate:
    canonical_url: str
    posted_date: date | None = None
    issue_date: date | None = None
    company_name: str | None = None
    country: str | None = None
    subject: str | None = None
    issuing_office: str | None = None
    marcs_cms_number: str | None = None
    fda_reference_number: str | None = None
    response_url: str | None = None
    closeout_url: str | None = None
    listing_ordinal: int = 0

    def to_ingestion_metadata(self, source_url: str) -> CandidateMetadata:
        return CandidateMetadata(
            canonical_url=self.canonical_url,
            posted_date=self.posted_date,
            issue_date=self.issue_date,
            company_name=self.company_name,
            country=self.country,
            subject=self.subject,
            issuing_office=self.issuing_office,
            marcs_cms_number=self.marcs_cms_number,
            fda_reference_number=self.fda_reference_number,
            http_provenance={
                "discovery_source_url": source_url,
                "listing_ordinal": self.listing_ordinal,
                "response_url": self.response_url,
                "closeout_url": self.closeout_url,
            },
        )


@dataclass(frozen=True)
class ListingRepresentation:
    source_url: str
    final_url: str
    media_type: str
    raw_sha256: str
    detected_columns: list[str]
    schema_fingerprint: str
    candidates: list[ListingCandidate]
    export_url: str | None = None
    warnings: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class DataTablesConfiguration:
    """Bounded server-side table configuration published by the FDA page."""

    ajax_url: str
    static_params: dict[str, str]
    columns: list[str]
    total_items: int

    def page_url(
        self,
        *,
        start: int,
        length: int,
        draw: int,
        allowed_hosts: list[str],
    ) -> str:
        params = dict(self.static_params)
        params.update(
            {
                "draw": str(draw),
                "start": str(max(0, start)),
                "length": str(max(1, min(length, 1000))),
                "search[value]": "",
                "search[regex]": "false",
                "order[0][column]": "0",
                "order[0][dir]": "desc",
            }
        )
        for index in range(len(self.columns)):
            prefix = f"columns[{index}]"
            params[f"{prefix}[data]"] = str(index)
            params[f"{prefix}[name]"] = ""
            params[f"{prefix}[searchable]"] = "true"
            params[f"{prefix}[orderable]"] = "false" if index == 7 else "true"
            params[f"{prefix}[search][value]"] = ""
            params[f"{prefix}[search][regex]"] = "false"
        separator = "&" if "?" in self.ajax_url else "?"
        return canonicalize_fda_url(f"{self.ajax_url}{separator}{urlencode(params)}", allowed_hosts)


def discover_datatables_configuration(
    html: bytes | str,
    *,
    page_url: str,
    allowed_hosts: list[str],
) -> DataTablesConfiguration | None:
    """Read FDA's public Drupal DataTables settings without executing page scripts."""

    soup = BeautifulSoup(html, "lxml")
    for script in soup.select("script"):
        raw = script.string or script.get_text() or ""
        raw = raw.strip()
        if not raw.startswith("{") or '"datatables"' not in raw:
            continue
        try:
            settings = json.loads(raw)
        except (TypeError, ValueError, json.JSONDecodeError):
            continue
        tables = settings.get("datatables")
        if not isinstance(tables, dict):
            continue
        for config in tables.values():
            if not isinstance(config, dict) or not config.get("serverSide"):
                continue
            ajax = config.get("ajax")
            if not isinstance(ajax, dict) or not ajax.get("url"):
                continue
            try:
                ajax_url = canonicalize_fda_url(urljoin(page_url, str(ajax["url"])), allowed_hosts)
            except UnsafeUrlError:
                continue
            raw_params = ajax.get("data")
            if not isinstance(raw_params, dict):
                continue
            static_params = {
                str(key): str(value).lower() if isinstance(value, bool) else str(value)
                for key, value in raw_params.items()
                if isinstance(value, (str, int, float, bool))
            }
            headers = config.get("aoColumnHeaders")
            columns = (
                [
                    normalize_whitespace(str(item.get("content") or ""))
                    for item in headers
                    if isinstance(item, dict)
                ]
                if isinstance(headers, list)
                else []
            )
            columns = [item for item in columns if item]
            if not columns:
                columns = [
                    normalize_whitespace(cell.get_text(" ", strip=True))
                    for cell in soup.select("table thead th")
                ]
            try:
                total_items = int(
                    config.get("deferLoading") or static_params.get("total_items") or 0
                )
            except (TypeError, ValueError):
                total_items = 0
            if not columns or total_items <= 0:
                continue
            return DataTablesConfiguration(
                ajax_url=ajax_url,
                static_params=static_params,
                columns=columns,
                total_items=min(total_items, 10_000),
            )
    return None


def discover_export_url(html: bytes | str, page_url: str, allowed_hosts: list[str]) -> str | None:
    soup = BeautifulSoup(html, "lxml")
    scored: list[tuple[int, str]] = []
    for link in soup.select("a[href]"):
        href = str(link.get("href", ""))
        label = normalize_whitespace(link.get_text(" ", strip=True)).casefold()
        candidate = urljoin(page_url, href)
        score = 0
        path = candidate.casefold()
        if ".xlsx" in path:
            score += 4
        if ".csv" in path:
            score += 3
        if "export" in label or "download" in label:
            score += 2
        if "warning letter" in label:
            score += 1
        if not score:
            continue
        try:
            canonical = canonicalize_fda_url(candidate, allowed_hosts)
        except UnsafeUrlError:
            continue
        scored.append((score, canonical))
    return max(scored, default=(0, None), key=lambda item: item[0])[1]


def _safe_url(value: object, base_url: str, allowed_hosts: list[str]) -> str | None:
    text = normalize_whitespace(str(value or ""))
    if not text or text.casefold() in {"n/a", "none", "no"}:
        return None
    try:
        return canonicalize_fda_url(urljoin(base_url, text), allowed_hosts)
    except UnsafeUrlError:
        return None


def _rows_to_candidates(
    rows: list[dict[str, Any]], source_url: str, allowed_hosts: list[str]
) -> list[ListingCandidate]:
    candidates: list[ListingCandidate] = []
    seen: set[str] = set()
    for ordinal, row in enumerate(rows):
        normalized = {_canonical_column(key): value for key, value in row.items()}
        url = _safe_url(normalized.get("canonical_url"), source_url, allowed_hosts)
        if not url or url in seen:
            continue
        seen.add(url)

        def optional_text(key: str, values: dict[str, Any] = normalized) -> str | None:
            value = normalize_whitespace(str(values.get(key) or ""))
            return value or None

        candidates.append(
            ListingCandidate(
                canonical_url=url,
                posted_date=_parse_date(normalized.get("posted_date")),
                issue_date=_parse_date(normalized.get("issue_date")),
                company_name=optional_text("company_name"),
                country=optional_text("country"),
                subject=optional_text("subject"),
                issuing_office=optional_text("issuing_office"),
                marcs_cms_number=optional_text("marcs_cms_number"),
                fda_reference_number=optional_text("fda_reference_number"),
                response_url=_safe_url(normalized.get("response_url"), source_url, allowed_hosts),
                closeout_url=_safe_url(normalized.get("closeout_url"), source_url, allowed_hosts),
                listing_ordinal=ordinal,
            )
        )
    return candidates


def _csv_rows(content: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    columns = [str(item) for item in (reader.fieldnames or [])]
    return columns, list(reader)


def _xlsx_rows(content: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    try:
        # FDA's export commonly stores the canonical letter URL as the Company
        # Name cell hyperlink rather than as a visible URL column. Normal mode
        # is required because openpyxl does not expose hyperlinks in read-only
        # worksheets. The acquisition byte limit bounds workbook size upstream.
        workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=True)
    except Exception as exc:
        raise ListingDiscoveryError("FDA XLSX export could not be opened") from exc
    worksheet = workbook.active
    iterator = worksheet.iter_rows()
    try:
        header = next(iterator)
    except StopIteration as exc:
        raise ListingDiscoveryError("FDA XLSX export is empty") from exc
    columns = [normalize_whitespace(str(cell.value or "")) for cell in header]
    rows: list[dict[str, Any]] = []
    for cells in iterator:
        values = [cell.value for cell in cells]
        row = dict(zip(columns, values, strict=False))
        for index, cell in enumerate(cells):
            if index >= len(columns) or not cell.hyperlink:
                continue
            target = str(cell.hyperlink.target or "")
            column = _canonical_column(columns[index])
            if column == "company_name":
                row["Letter URL"] = target
            elif column == "response_url":
                row["Response URL"] = target
            elif column == "closeout_url":
                row["Closeout URL"] = target
        rows.append(row)
    workbook.close()
    return columns, rows


def _html_rows(
    content: bytes, source_url: str, allowed_hosts: list[str]
) -> tuple[list[str], list[dict[str, Any]], str | None]:
    soup = BeautifulSoup(content, "lxml")
    export_url = discover_export_url(content, source_url, allowed_hosts)
    table = soup.select_one("table")
    if table is None:
        return [], [], export_url
    headers = [
        normalize_whitespace(cell.get_text(" ", strip=True)) for cell in table.select("thead th")
    ]
    body_rows = table.select("tbody tr") or table.select("tr")[1:]
    if not headers:
        first = table.select_one("tr")
        headers = (
            [normalize_whitespace(cell.get_text(" ", strip=True)) for cell in first.select("th,td")]
            if first
            else []
        )
    rows: list[dict[str, Any]] = []
    for row in body_rows:
        cells = row.select("th,td")
        values: list[str] = []
        for cell in cells:
            link = cell.select_one("a[href]")
            header_index = len(values)
            header = _canonical_column(headers[header_index]) if header_index < len(headers) else ""
            if link and header in {"canonical_url", "response_url", "closeout_url"}:
                values.append(str(link.get("href", "")))
            elif link and header == "company_name":
                values.append(normalize_whitespace(link.get_text(" ", strip=True)))
            else:
                values.append(normalize_whitespace(cell.get_text(" ", strip=True)))
        mapped = dict(zip(headers, values, strict=False))
        # The canonical detail URL is commonly attached to the company cell.
        if not any(_canonical_column(key) == "canonical_url" for key in mapped):
            company_cell = next(
                (
                    cell
                    for index, cell in enumerate(cells)
                    if index < len(headers) and _canonical_column(headers[index]) == "company_name"
                ),
                None,
            )
            company_link = company_cell.select_one("a[href]") if company_cell else None
            if company_link:
                mapped["Letter URL"] = str(company_link.get("href", ""))
        rows.append(mapped)
    return headers, rows, export_url


def parse_datatables_page(
    content: bytes,
    *,
    source_url: str,
    columns: list[str],
    allowed_hosts: list[str],
) -> ListingRepresentation:
    """Parse one official FDA server-side DataTables JSON response."""

    try:
        payload = json.loads(content)
    except (UnicodeDecodeError, TypeError, ValueError, json.JSONDecodeError) as exc:
        raise ListingDiscoveryError("FDA DataTables response was not valid JSON") from exc
    data = payload.get("data") if isinstance(payload, dict) else None
    if not isinstance(data, list):
        raise ListingDiscoveryError("FDA DataTables response omitted its data rows")

    rows: list[dict[str, Any]] = []
    for values in data:
        if not isinstance(values, list):
            continue
        mapped: dict[str, Any] = {}
        for index, column in enumerate(columns):
            if index >= len(values):
                break
            fragment = BeautifulSoup(str(values[index] or ""), "lxml")
            mapped[column] = normalize_whitespace(fragment.get_text(" ", strip=True))
            link = fragment.select_one("a[href]")
            if not link:
                continue
            canonical = _canonical_column(column)
            if canonical == "company_name":
                mapped["Letter URL"] = str(link.get("href", ""))
            elif canonical == "response_url":
                mapped["Response URL"] = str(link.get("href", ""))
            elif canonical == "closeout_url":
                mapped["Closeout URL"] = str(link.get("href", ""))
        rows.append(mapped)

    canonical_columns = [_canonical_column(column) for column in columns if _column_key(column)]
    candidates = _rows_to_candidates(rows, source_url, allowed_hosts)
    if rows and not candidates:
        raise ListingDiscoveryError(
            "FDA DataTables response contained rows but no valid warning-letter URLs"
        )
    final = canonicalize_fda_url(source_url, allowed_hosts)
    fingerprint = hashlib.sha256(
        json.dumps(canonical_columns, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    warnings = ["zero_listing_rows"] if not rows else []
    return ListingRepresentation(
        source_url=final,
        final_url=final,
        media_type="json",
        raw_sha256=hashlib.sha256(content).hexdigest(),
        detected_columns=canonical_columns,
        schema_fingerprint=fingerprint,
        candidates=candidates,
        warnings=warnings,
    )


def parse_listing(
    content: bytes,
    *,
    source_url: str,
    final_url: str | None = None,
    content_type: str | None = None,
    allowed_hosts: list[str],
) -> ListingRepresentation:
    final = canonicalize_fda_url(final_url or source_url, allowed_hosts)
    media = (content_type or "").split(";", 1)[0].casefold()
    if (
        content.startswith(b"PK\x03\x04")
        or "spreadsheet" in media
        or final.casefold().endswith(".xlsx")
    ):
        format_name = "xlsx"
        columns, rows = _xlsx_rows(content)
        export_url = final
    elif media == "text/csv" or final.casefold().endswith(".csv"):
        format_name = "csv"
        columns, rows = _csv_rows(content)
        export_url = final
    else:
        format_name = "html"
        columns, rows, export_url = _html_rows(content, final, allowed_hosts)

    canonical_columns = [_canonical_column(column) for column in columns if _column_key(column)]
    fingerprint = hashlib.sha256(
        json.dumps(canonical_columns, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    candidates = _rows_to_candidates(rows, final, allowed_hosts)
    warnings: list[str] = []
    if "canonical_url" not in canonical_columns and not candidates:
        warnings.append("missing_canonical_url_column")
    if not rows:
        warnings.append("zero_listing_rows")
    if rows and not candidates:
        raise ListingDiscoveryError(
            "FDA listing contained rows but no valid allowlisted warning-letter URLs"
        )
    return ListingRepresentation(
        source_url=canonicalize_fda_url(source_url, allowed_hosts),
        final_url=final,
        media_type=format_name,
        raw_sha256=hashlib.sha256(content).hexdigest(),
        detected_columns=canonical_columns,
        schema_fingerprint=fingerprint,
        candidates=candidates,
        export_url=export_url,
        warnings=warnings,
    )
